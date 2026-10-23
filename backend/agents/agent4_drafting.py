"""
agent4_drafting.py — LawGrid drafting agent

Uses Groq when configured. Without a key (or when the call fails) it generates
review-ready drafts from local templates so the core pipeline still produces
documents. Every document is explicitly marked as a draft for human review.
"""

from __future__ import annotations

import json
import logging
import os
import re
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

from dotenv import load_dotenv
from pydantic import ValidationError

from case_state import GeneratedDoc, GeneratedDocuments

load_dotenv(override=False)
logger = logging.getLogger(__name__)

GROQ_MODEL = "llama-3.3-70b-versatile"
TEMPLATE_DIR = Path(__file__).resolve().parent / "templates"
OUTPUT_DIR = Path(__file__).resolve().parent.parent / "output_docs"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

DISCLAIMER = (
    "\n\n---\n**DISCLAIMER:** This document is a draft generated with automated "
    "tools. It is not legal advice. A qualified advocate should verify the facts, "
    "sections, limitation periods, court fees and formatting before filing."
)


def _get_client():
    key = os.environ.get("GROQ_API_KEY")
    if not key:
        return None
    try:
        from groq import Groq

        return Groq(api_key=key)
    except Exception as exc:  # pragma: no cover - depends on optional SDK
        logger.warning("Could not initialise Groq client: %s", exc)
        return None


def _load_template(doc_type: str) -> str:
    template_map = {
        "fir": "fir.txt",
        "legal_notice": "legal_notice.txt",
        "consumer_complaint": "consumer_complaint.txt",
    }
    filename = template_map.get(doc_type)
    if not filename:
        return ""
    path = TEMPLATE_DIR / filename
    if path.exists():
        return path.read_text(encoding="utf-8")
    return ""


def _call_with_retry(
    prompt: str,
    system_prompt: str,
    client,
    max_attempts: int = 2,
) -> dict:
    last_error = None
    for attempt in range(1, max_attempts + 1):
        try:
            response = client.chat.completions.create(
                model=GROQ_MODEL,
                temperature=0.0,
                response_format={"type": "json_object"},
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": prompt},
                ],
            )
            raw = response.choices[0].message.content.strip()
            if raw.startswith("```"):
                raw = raw.split("```")[1]
                if raw.lstrip().startswith("json"):
                    raw = raw.lstrip()[4:]
                raw = raw.strip()
            return json.loads(raw)
        except Exception as exc:
            last_error = exc
            wait = 2 ** attempt
            logger.warning(
                "[agent4] attempt %s failed: %s - retrying in %ss",
                attempt,
                exc,
                wait,
            )
            time.sleep(wait)
    raise RuntimeError(f"Groq failed after {max_attempts} attempts: {last_error}")


def _safe_filename(value: str) -> str:
    clean = re.sub(r"[^A-Za-z0-9._-]+", "_", value or "case")
    return clean.strip("._")[:120] or "case"


def _build_context(case_state: dict) -> dict:
    facts = case_state.get("structured_facts") or {}
    mapping = case_state.get("legal_mapping") or {}
    plan = case_state.get("action_plan") or {}
    evidence = case_state.get("evidence_inventory") or []

    parties = facts.get("parties") or []
    complainant = next(
        (
            party.get("name") or party.get("description") or "Complainant"
            for party in parties
            if party.get("role") == "complainant"
        ),
        "Complainant",
    )
    respondent = next(
        (
            party.get("name") or party.get("description") or "Respondent"
            for party in parties
            if party.get("role") == "respondent"
        ),
        "Respondent",
    )

    all_sections = (mapping.get("applicable_sections") or []) + (
        mapping.get("bns_bnss_sections") or []
    )
    sections_block = "\n".join(
        f"- {section.get('section_ref', 'Section')} - "
        f"{section.get('act_name', 'Act')}: "
        f"{section.get('description') or 'Description not available.'}"
        for section in all_sections
    ) or "- Applicable provisions to be verified by an advocate."

    facts_list = [str(fact) for fact in (facts.get("key_facts") or []) if fact]
    facts_block = (
        "\n".join(f"{index + 1}. {fact}" for index, fact in enumerate(facts_list))
        or "1. The facts are to be completed from the client's instructions."
    )

    evidence_block = "\n".join(
        f"- {item.get('evidence_type', 'Document')} "
        f"({item.get('file_reference') or 'reference'}): "
        f"{item.get('description') or ''}".rstrip(": ")
        for item in evidence
    ) or "- No documents have been attached yet."

    return {
        "complainant": complainant,
        "respondent": respondent,
        "dispute_type": str(
            facts.get("incident_type") or "legal dispute"
        ).replace("_", " "),
        "case_title": facts.get("case_title") or "Legal Dispute",
        "summary": facts.get("incident_summary") or "",
        "incident_date": facts.get("incident_date") or "",
        "state": case_state.get("state_jurisdiction") or "India",
        "date": datetime.now().strftime("%B %d, %Y"),
        "facts": facts_list,
        "facts_block": facts_block,
        "sections_block": sections_block,
        "forum": plan.get("forum_selection") or "appropriate forum",
        "demand_amount": plan.get("demand_amount") or "As per actual loss and damages",
        "relief_sought": plan.get("relief_sought") or "Appropriate relief under law",
        "escalation_path": plan.get("escalation_path") or "",
        "evidence_block": evidence_block,
        "evidence_checklist": plan.get("evidence_checklist") or [],
        "language": case_state.get("language_preference") or "english",
    }


def _header(title: str, context: dict) -> str:
    return (
        f"# {title}\n\n"
        f"**Prepared for:** {context['complainant']}\n\n"
        f"**Date:** {context['date']}\n\n"
        f"**Jurisdiction:** {context['state']}\n\n"
        "> **DRAFT FOR REVIEW.** This template-based draft must be checked and "
        "adapted by a qualified advocate before it is signed, sent or filed.\n\n"
    )


def _fallback_document(doc_type: str, context: dict) -> tuple[str, str]:
    """Generate a useful, clearly-labelled draft without any external API."""

    facts = context["facts_block"]
    sections = context["sections_block"]
    complainant = context["complainant"]
    respondent = context["respondent"]
    forum = context["forum"]
    relief = context["relief_sought"]
    amount = context["demand_amount"]
    state = context["state"]
    dispute = context["dispute_type"]
    summary = context["summary"] or "The facts are summarised in the numbered paragraphs below."

    if doc_type == "consumer_complaint":
        title = "Consumer Complaint"
        content = (
            _header(title, context)
            + "**BEFORE THE DISTRICT CONSUMER DISPUTES REDRESSAL COMMISSION**\n\n"
            f"**{state}**\n\n"
            "Complaint No. ________ of 2026\n\n"
            f"**Complainant:** {complainant}\n\n"
            f"**Opposite Party:** {respondent}\n\n"
            "## 1. Facts of the case\n\n"
            f"{facts}\n\n"
            f"{summary}\n\n"
            "## 2. Jurisdiction and limitation\n\n"
            "The complaint is filed under the Consumer Protection Act, 2019. "
            "The complainant is a 'consumer' and the opposite party has provided "
            "deficient services / engaged in an unfair trade practice. The claim "
            "is within the pecuniary jurisdiction of this Commission. The cause "
            "of action arose on the date stated in the facts above.\n\n"
            "## 3. Legal grounds\n\n"
            f"{sections}\n\n"
            "## 4. Documents relied upon\n\n"
            f"{context['evidence_block']}\n\n"
            "## 5. Prayer\n\n"
            "The complainant respectfully prays that this Commission may be pleased to:\n\n"
            f"1. Direct the opposite party to {relief.lower()};\n"
            f"2. Award compensation and litigation costs;\n"
            "3. Pass any other order that this Commission considers just and proper.\n\n"
            "**Verification:** I, the complainant, verify that the facts stated above "
            "are true and correct to the best of my knowledge and belief.\n\n"
            f"Place: {state}\n\nDate: {context['date']}\n\n"
            f"Signature: ____________________\n\n{complainant}\n"
            + DISCLAIMER
        )
    elif doc_type == "fir_draft":
        title = "Complaint for Registration of FIR"
        content = (
            _header(title, context)
            + "**To,**\n\nThe Station House Officer,\n\n"
            f"[Police Station], {state}\n\n"
            f"**Subject:** Information regarding a cognizable offence and request "
            f"for registration of an FIR under Section 173 of the Bharatiya Nagarik "
            f"Suraksha Sanhita, 2023.\n\n"
            f"Respected Sir/Madam,\n\n"
            f"I, {complainant}, wish to report the following incident:\n\n"
            f"{facts}\n\n"
            f"{summary}\n\n"
            "The above facts disclose a cognizable offence. I request that the "
            "information be recorded and an FIR registered. I am ready to provide "
            "all documents and to cooperate with the investigation.\n\n"
            "**Applicable provisions (to be verified):**\n\n"
            f"{sections}\n\n"
            "**Documents enclosed:**\n\n"
            f"{context['evidence_block']}\n\n"
            "**Requested action:** Registration of an FIR, preservation of CCTV "
            "and other electronic evidence, and a written acknowledgement of this "
            "complaint.\n\n"
            f"Place: {state}\n\nDate: {context['date']}\n\n"
            f"Signature: ____________________\n\n{complainant}\n"
            + DISCLAIMER
        )
    elif doc_type == "rti_application":
        title = "Application under the Right to Information Act, 2005"
        content = (
            _header(title, context)
            + "**To,**\n\nThe Public Information Officer,\n\n"
            f"[Public Authority], {state}\n\n"
            "**Subject:** Request for information under Section 6 of the Right to "
            "Information Act, 2005.\n\n"
            "1. **Name of applicant:** " + complainant + "\n\n"
            "2. **Address:** [Complete postal address]\n\n"
            "3. **Contact details:** [Phone / email]\n\n"
            "4. **Information sought:** Please provide the following information, "
            "including certified copies of the relevant records:\n\n"
            f"{facts}\n\n"
            "5. **Period to which the information relates:** [Specify dates]\n\n"
            "6. **Whether the applicant is a citizen of India:** Yes\n\n"
            "7. **Application fee:** [Amount and mode of payment]\n\n"
            "If the information is held by another public authority, please transfer "
            "this application under Section 6(3) and inform me in writing. If any "
            "part is exempt, please provide the reasons and the gist of the exempt "
            "material as required by law.\n\n"
            f"Place: {state}\n\nDate: {context['date']}\n\n"
            f"Signature: ____________________\n\n{complainant}\n"
            + DISCLAIMER
        )
    elif doc_type == "affidavit":
        title = "Affidavit"
        content = (
            _header(title, context)
            + "**IN THE COURT OF COMPETENT JURISDICTION**\n\n"
            f"**{state}**\n\n"
            "AFFIDAVIT OF " + str(complainant).upper() + "\n\n"
            f"I, {complainant}, aged about ____ years, resident of "
            "[complete address], do hereby solemnly affirm and state on oath as "
            "follows:\n\n"
            f"{facts}\n\n"
            f"{summary}\n\n"
            "I state that the facts stated above are true and correct to the best "
            "of my knowledge and belief. No material fact has been concealed. I "
            "am competent to swear this affidavit.\n\n"
            "**Documents referred to:**\n\n"
            f"{context['evidence_block']}\n\n"
            f"Place: {state}\n\nDate: {context['date']}\n\n"
            "Deponent: ____________________\n\n"
            "**Verification:** Verified at " + state + " on this ____ day of "
            "________, 2026 that the contents of the above affidavit are true and "
            "correct.\n\n"
            "Deponent\n"
            + DISCLAIMER
        )
    elif doc_type == "corruption_complaint":
        title = "Complaint regarding Corruption / Abuse of Public Office"
        content = (
            _header(title, context)
            + "**To,**\n\nThe Director / Competent Authority,\n\n"
            "Central Vigilance Commission / State Lokayukta\n\n"
            f"**Subject:** Complaint regarding alleged corruption by a public "
            f"servant in {state}.\n\n"
            f"I, {complainant}, wish to bring the following facts to your notice:\n\n"
            f"{facts}\n\n"
            f"{summary}\n\n"
            "**Relevant legal provisions (to be verified):**\n\n"
            f"{sections}\n\n"
            "I request an inquiry into the conduct of the public servant named in "
            "the facts above. I am willing to provide further information and "
            "documents, and request that my identity be protected to the extent "
            "permitted by law.\n\n"
            "**Documents enclosed:**\n\n"
            f"{context['evidence_block']}\n\n"
            f"Place: {state}\n\nDate: {context['date']}\n\n"
            f"Signature: ____________________\n\n{complainant}\n"
            + DISCLAIMER
        )
    elif doc_type == "lawyer_brief":
        title = "Case Brief for Counsel"
        content = (
            _header(title, context)
            + "**CONFIDENTIAL - PREPARED FOR COUNSEL**\n\n"
            f"**Client:** {complainant}\n\n"
            f"**Opposite party:** {respondent}\n\n"
            f"**Dispute:** {dispute}\n\n"
            f"**Forum under consideration:** {forum}\n\n"
            "## 1. Executive summary\n\n"
            f"{summary}\n\n"
            "## 2. Chronology and key facts\n\n"
            f"{facts}\n\n"
            "## 3. Legal provisions identified by the research layer\n\n"
            f"{sections}\n\n"
            "## 4. Evidence available / required\n\n"
            f"{context['evidence_block']}\n\n"
            "**Evidence checklist:**\n\n"
            + "\n".join(
                f"- {item}" for item in context["evidence_checklist"]
            )
            + "\n\n## 5. Proposed strategy\n\n"
            f"- Forum: {forum}\n"
            f"- Relief sought: {relief}\n"
            f"- Demand: {amount}\n"
            f"- Escalation path: {context['escalation_path'] or 'To be advised'}\n\n"
            "## 6. Issues for counsel to verify\n\n"
            "- Limitation and cause of action.\n"
            "- Correct forum and pecuniary jurisdiction.\n"
            "- Whether any interim relief is required.\n"
            "- Accuracy of every statutory reference in this brief.\n"
            + DISCLAIMER
        )
    else:
        title = "Legal Notice"
        content = (
            _header(title, context)
            + f"**To:**\n\n{respondent}\n\n"
            f"**From:**\n\n{complainant}\n\n"
            f"**Subject:** Legal notice regarding {dispute}.\n\n"
            "I, the undersigned, am constrained to issue this legal notice to you "
            "in connection with the following facts:\n\n"
            f"{facts}\n\n"
            f"{summary}\n\n"
            "**Legal basis (to be verified by an advocate):**\n\n"
            f"{sections}\n\n"
            "Despite the above, you have failed to resolve the dispute. I therefore "
            "call upon you to comply with the following demand within **fifteen (15) "
            "days** of receipt of this notice:\n\n"
            f"- {relief}\n"
            f"- Payment/compensation: {amount}\n"
            f"- Any other relief required to make good the loss: {context['escalation_path'] or 'as advised'}\n\n"
            "If you fail to comply, I will be constrained to initiate appropriate "
            "civil, consumer, criminal or regulatory proceedings against you at your "
            "risk as to costs and consequences. This notice is issued without "
            "prejudice to all rights and remedies available in law.\n\n"
            "A copy of this notice is retained for record and may be sent by "
            "registered post, courier and email.\n\n"
            f"Place: {state}\n\nDate: {context['date']}\n\n"
            f"Signature: ____________________\n\n{complainant}\n"
            + DISCLAIMER
        )

    return title, content


def _export_to_docx(content_md: str, filepath: Path) -> bool:
    try:
        import docx

        document = docx.Document()
        for line in content_md.split("\n"):
            clean = re.sub(r"^#{1,6}\s*", "", line)
            clean = clean.replace("**", "").replace("__", "")
            if clean.strip() == "---":
                continue
            document.add_paragraph(clean)
        document.save(str(filepath))
        return True
    except Exception as exc:
        logger.warning("DOCX export failed: %s", exc)
        return False


def _export_to_pdf(content_md: str, filepath: Path) -> bool:
    try:
        from markdown_pdf import MarkdownPdf, Section

        pdf = MarkdownPdf()
        pdf.add_section(Section(content_md, toc=False))
        pdf.save(str(filepath))
        return True
    except Exception as exc:
        logger.warning("PDF export failed: %s", exc)
        return False


def _call_groq(prompt: str, client) -> dict:
    system_prompt = """You are an expert Indian legal drafting assistant.
Draft a formal, filing-ready document in Markdown. Output ONLY valid JSON:
{"content_md": "the complete draft", "document_title": "formal title"}

Rules:
- Use only the facts and statutory sections provided. Never invent section numbers.
- Include placeholders such as [Court], [Police Station] where facts are missing.
- Use formal Indian legal English. Keep the structure clear and numbered.
- Include a verification/signature block and a disclaimer.
- Do not claim to be a lawyer or give a guarantee of outcome."""
    return _call_with_retry(prompt, system_prompt, client)


def drafting_agent(case_state: dict) -> dict:
    start = time.time()

    plan = case_state.get("action_plan") or {}
    doc_types = list(plan.get("document_types_required") or ["legal_notice"])
    if plan.get("lawyer_recommended") and "lawyer_brief" not in doc_types:
        doc_types.append("lawyer_brief")

    context = _build_context(case_state)
    client = _get_client()
    generated: dict[str, dict[str, Any]] = {}

    for doc_type in doc_types:
        template_content = _load_template(doc_type)
        template_instruction = (
            f"\nUse this template as a structural reference:\n{template_content}\n"
            if template_content
            else ""
        )
        prompt = f"""
Generate a professional {doc_type} for this case.
{template_instruction}
Case details:
- Dispute type: {context['dispute_type']}
- Complainant: {context['complainant']}
- Respondent: {context['respondent']}
- State: {context['state']}
- Forum: {context['forum']}
- Facts:
{context['facts_block']}
- Law sections:
{context['sections_block']}
- Evidence:
{context['evidence_block']}
- Relief: {context['relief_sought']}
- Demand: {context['demand_amount']}

Return only the requested JSON schema.
"""

        result: Optional[dict] = None
        fallback = False
        if client is not None:
            try:
                result = _call_groq(prompt, client)
            except Exception as exc:
                logger.warning("[agent4] generation failed: %s - using template", exc)
                case_state.setdefault("human_review_flags", []).append(
                    f"Groq drafting unavailable ({exc}); template draft used for {doc_type}"
                )
                fallback = True
        else:
            fallback = True

        if not result or not result.get("content_md"):
            title, content = _fallback_document(doc_type, context)
            result = {"content_md": content, "document_title": title}
            fallback = True

        content = str(result.get("content_md") or "").strip()
        if "DISCLAIMER" not in content.upper():
            content += DISCLAIMER

        safe_title = _safe_filename(
            f"{case_state.get('case_id', 'case')}_{doc_type}"
        )
        docx_path = OUTPUT_DIR / f"{safe_title}.docx"
        pdf_path = OUTPUT_DIR / f"{safe_title}.pdf"
        docx_ok = _export_to_docx(content, docx_path)
        pdf_ok = _export_to_pdf(content, pdf_path)

        generated[doc_type] = {
            "content_md": content,
            "document_title": result.get("document_title") or doc_type.replace("_", " ").title(),
            "docx_url": str(docx_path) if docx_ok else None,
            "pdf_url": str(pdf_path) if pdf_ok else None,
            "disclaimer_included": True,
            "is_fallback": fallback,
        }

    existing_docs = case_state.get("generated_documents")
    if existing_docs is None:
        existing_docs = GeneratedDocuments()
    elif isinstance(existing_docs, dict):
        existing_docs = GeneratedDocuments(**existing_docs)

    for doc_type, doc_data in generated.items():
        if hasattr(existing_docs, doc_type):
            setattr(
                existing_docs,
                doc_type,
                GeneratedDoc(
                    content_md=doc_data["content_md"],
                    docx_url=doc_data["docx_url"],
                    pdf_url=doc_data["pdf_url"],
                    disclaimer_included=doc_data["disclaimer_included"],
                ),
            )

    case_state["generated_documents"] = existing_docs.model_dump()

    duration_ms = int((time.time() - start) * 1000)
    now = datetime.now(timezone.utc).isoformat()
    case_state.setdefault("agent_trace", []).append(
        {
            "agent": "drafting",
            "timestamp": now,
            "duration_ms": duration_ms,
            "docs_generated": list(generated.keys()),
            "mode": "template" if all(d.get("is_fallback") for d in generated.values()) else "groq_or_mixed",
        }
    )
    case_state.setdefault("case_timeline", []).append(
        {"event": "Drafting agent completed", "timestamp": now}
    )

    return case_state
