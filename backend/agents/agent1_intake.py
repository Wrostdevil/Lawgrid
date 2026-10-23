from __future__ import annotations

import io
import json
import logging
import os
import re
import tempfile
import time
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Optional

from groq import Groq
from sarvamai import SarvamAI

from case_state import CaseState, EvidenceItem, OcrDocument, Party, StructuredFacts

logger = logging.getLogger(__name__)


GROQ_MODEL = "llama-3.3-70b-versatile"
MAX_TOKENS = 2048
MAX_FOLLOWUP_TOKENS = 200


def _call_with_retry(prompt: str, system_prompt: str, client: Groq, max_attempts: int = 3, json_mode: bool = True) -> dict:
    """Call Groq with exponential backoff and JSON validation."""
    last_error = None
    
    for attempt in range(1, max_attempts + 1):
        try:
            extra_args = {"response_format": {"type": "json_object"}} if json_mode else {}
            
            response = client.chat.completions.create(
                model=GROQ_MODEL,
                temperature=0.0,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": prompt}
                ],
                **extra_args
            )
            raw = response.choices[0].message.content.strip()
            
            # Robust JSON extraction
            start_idx = raw.find('{')
            end_idx = raw.rfind('}') + 1
            if start_idx != -1 and end_idx != 0:
                raw_json = raw[start_idx:end_idx]
                return json.loads(raw_json)
            else:
                raise ValueError("No JSON found in response")
                
        except Exception as e:
            last_error = e
            wait = 2 ** attempt
            logger.warning(f"[groq] attempt {attempt} failed: {e} — retrying in {wait}s")
            time.sleep(wait)
            
    raise RuntimeError(f"Groq failed after {max_attempts} attempts: {last_error}")


INTAKE_SYSTEM_PROMPT = """You are a Professional Legal Intake Officer at a high-end Indian law firm.

Your job is to analyse a citizen's raw narrative and extract a precise, structured incident report. 

TONE & BEHAVIOR:
1. PERSONA: Professional Legal Intake Clerk. Direct and efficient.
2. DELTA-ONLY: The `new_facts` array MUST strictly contain only information provided in the VERY LAST user message.
3. ABSOLUTELY NO RECAP: Do NOT include facts in `new_facts` that were already mentioned in the prior chat history. If the user repeats themselves, `new_facts` should be empty.
4. POINTED QUESTIONS: Focus `follow_up_query` on the highest priority missing data.
5. NO FLUFF: No "I have noted that" or "Acknowledging your facts" phrases.

Return ONLY valid JSON.

If the user is just saying hello, return "GREETING" in intent. If the input is nonsensical/gibberish, return "ABSURD" in intent. Otherwise return "LEGAL".

7. ALLOWED CATEGORIES: You MUST categorize the legal issue into one of these strict keys for the `incident_type` field:
   - `consumer` (Consumer disputes, product defects, services)
   - `tenant` (Rent, eviction, landlord issues)
   - `labour` (Employment, salary, harassment at work)
   - `criminal` (Theft, assault, FIRs, bail)
   - `cyber` (Online fraud, hacking, harassment)
   - `property` (Land disputes, inheritance, registration)
   - `family` (Marriage, custody, inheritance)
   - `divorce` (Separation, alimony)
   - `debt_recovery` (Loans, recovery, cheque bounce)
   - `civil` (General contracts, damages, non-categorized civil)
   - `other` (Use ONLY if none of the above apply)

Return ONLY valid JSON.
{
  "intent": "string (GREETING/ABSURD/LEGAL)",
  "new_facts": ["string (Only NEWly discovered facts from the current turn)"],
  "follow_up_query": "string (A professional direct question for missing information)",
  "conversational_response": "string (OPTIONAL: Only used for GREETING or ABSURD intents. For LEGAL, leave null as it will be constructed from new_facts and follow_up_query.)",
  "language_preference": "string (hindi/english/hinglish)",
  "state_jurisdiction": "string (e.g. Delhi, Maharashtra) or null",
  "structured_facts": {
    "incident_type": "string (Must be one of the ALLOWED CATEGORIES keys)",
        "case_title": "string (short case title in plain language, max 90 chars) or null",
    "incident_summary": "string or null",
    "incident_date": "string or null",
    "parties": [
      {
        "role": "string",
        "name": "string or null",
        "contact": "string or null",
        "description": "string or null"
      }
    ],
    "timeline": ["string"],
    "urgency_level": "string",
    "monetary_value_inr": "number or null",
    "key_facts": ["string"],
    "missing_information": ["string"],
    "dispute_context": "string or null"
  },
  "evidence_inventory": [
    {
      "file_reference": "string (e.g. invoice, medical report)",
      "evidence_type": "string (receipt/agreement/photo/other)",
      "description": "string (concise description)"
    }
  ]
}
"""


FOLLOWUP_SYSTEM_PROMPT = """You convert missing legal-intake details into natural follow-up questions.

Rules:
Return ONLY a JSON array of strings.
Ask at most 3 concise questions.
"""


# ─────────────────────────────────────────────
# Language mapping
# ─────────────────────────────────────────────

def _detect_language_code(language_preference: str) -> str:

    mapping = {
        "hindi": "hi-IN",
        "hinglish": "hi-IN",
        "english": "en-IN",
    }

    return mapping.get(language_preference, "en-IN")


# ─────────────────────────────────────────────
# Offline / local fallback intake
# ─────────────────────────────────────────────

INDIAN_STATES = [
    "Andhra Pradesh", "Arunachal Pradesh", "Assam", "Bihar", "Chhattisgarh",
    "Goa", "Gujarat", "Haryana", "Himachal Pradesh", "Jharkhand", "Karnataka",
    "Kerala", "Madhya Pradesh", "Maharashtra", "Manipur", "Meghalaya",
    "Mizoram", "Nagaland", "Odisha", "Punjab", "Rajasthan", "Sikkim",
    "Tamil Nadu", "Telangana", "Tripura", "Uttar Pradesh", "Uttarakhand",
    "West Bengal", "Delhi", "Jammu and Kashmir", "Ladakh", "Puducherry",
    "Chandigarh", "Andaman and Nicobar Islands", "Dadra and Nagar Haveli and Daman and Diu",
    "Lakshadweep",
]

_GREETING_RE = re.compile(
    r"^\s*(hi|hello|hey|namaste|namaskar|good\s+(morning|afternoon|evening))\b",
    re.IGNORECASE,
)

_INCIDENT_KEYWORDS = {
    "family_domestic_violence": [
        "domestic violence", "husband beat", "husband beats", "wife beat",
        "dowry", "in-laws", "in laws", "marital cruelty", "abuse at home",
    ],
    "family_maintenance": [
        "maintenance", "alimony", "child support", "divorce", "separation",
        "custody",
    ],
    "cyber_crime": [
        "cyber", "online fraud", "online scam", "otp", "phishing", "hacked",
        "hacking", "fake profile", "upi fraud", "account fraud", "internet fraud",
    ],
    "rti_refusal": [
        "rti", "right to information", "public information officer", "pio",
    ],
    "corruption_complaint": [
        "bribe", "bribery", "corruption", "lokayukta", "cvc", "undue advantage",
    ],
    "debt_recovery": [
        "cheque bounce", "cheque bounced", "dishonoured cheque", "loan recovery",
        "money recovery", "borrowed money", "failed to repay", "not repaying",
    ],
    "consumer_fraud": [
        "consumer", "defective", "defect", "warranty", "refund", "product",
        "seller", "shop", "purchase", "service deficiency", "replacement",
        "misleading advertisement",
    ],
    "tenant_landlord_deposit": [
        "rent", "landlord", "tenant", "eviction", "security deposit",
        "rent agreement", "lease", "vacate",
    ],
    "wage_theft": [
        "salary", "wages", "unpaid salary", "salary not paid", "deduction",
        "employer", "labour", "labor", "workplace",
    ],
    "wrongful_termination": [
        "termination", "terminated", "fired", "dismissed", "retrenchment",
        "lost my job", "job termination",
    ],
    "property_dispute": [
        "land", "flat", "builder", "registry", "possession", "partition",
        "inheritance", "property", "real estate", "rera",
    ],
    "criminal_assault": [
        "assault", "hit me", "beat me", "attack", "injury", "hurt me",
        "physical violence",
    ],
    "criminal_fraud": [
        "fraud", "cheating", "cheated", "scam", "theft", "stolen", "threat",
        "criminal intimidation", "fir", "police complaint",
    ],
}

_RESPONDENT_TERMS = [
    ("landlord", "Landlord"),
    ("employer", "Employer"),
    ("company", "Company / seller"),
    ("seller", "Seller"),
    ("shop", "Shop / seller"),
    ("builder", "Builder / developer"),
    ("bank", "Bank"),
    ("husband", "Husband"),
    ("wife", "Wife"),
    ("police", "Police"),
    ("officer", "Public officer"),
    ("tenant", "Tenant"),
    ("borrower", "Borrower"),
]

_HINGLISH_MARKERS = {
    "hai", "nahi", "mujhe", "mera", "meri", "kya", "kyun", "kyu", "kar",
    "karna", "hua", "huaa", "gaya", "gayi", "paise", "paisa", "landlord",
    "kiraya", "salary", "nahin", "bhai", "sahab", "bina", "liye", "wala",
    "wale", "raha", "rahi", "tha", "thi", "hain", "aap", "hum", "unhone",
}


def _last_user_turn(narrative: str) -> str:
    text = (narrative or "").strip()
    if "User:" in text:
        return text.rsplit("User:", 1)[1].strip()
    return text


def _detect_language(text: str, fallback: str = "english") -> str:
    if re.search(r"[\u0900-\u097F]", text or ""):
        return "hindi"
    words = set(re.findall(r"[a-zA-Z]+", (text or "").lower()))
    if len(words.intersection(_HINGLISH_MARKERS)) >= 2:
        return "hinglish"
    return fallback if fallback in {"hindi", "english", "hinglish"} else "english"


def _detect_state(text: str) -> Optional[str]:
    for state in sorted(INDIAN_STATES, key=len, reverse=True):
        if re.search(rf"\b{re.escape(state)}\b", text or "", re.IGNORECASE):
            return state
    return None


def _extract_date(text: str) -> Optional[str]:
    patterns = [
        r"\b\d{1,2}[./-]\d{1,2}[./-]\d{2,4}\b",
        r"\b\d{1,2}(?:st|nd|rd|th)?\s+(?:January|February|March|April|May|June|July|August|September|October|November|December),?\s+\d{4}\b",
        r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)"
        r"\s+\d{1,2}(?:st|nd|rd|th)?,?\s+\d{4}\b",
        r"\b(?:today|yesterday|last week|last month|this week|this month)\b",
    ]
    for pattern in patterns:
        match = re.search(pattern, text or "", re.IGNORECASE)
        if match:
            return match.group(0)
    return None


def _extract_amount(text: str) -> Optional[float]:
    match = re.search(
        r"(?:₹|Rs\.?|INR)\s*([\d,]+(?:\.\d{1,2})?)",
        text or "",
        re.IGNORECASE,
    )
    if not match:
        return None
    try:
        return float(match.group(1).replace(",", ""))
    except ValueError:
        return None


def _detect_mentioned_evidence(text: str) -> list[dict]:
    """Identify evidence the citizen says they already have."""
    lowered = (text or "").lower()
    candidates = [
        ("invoice", "invoice", "Invoice"),
        ("bill", "bill", "Bill / receipt"),
        ("receipt", "receipt", "Receipt"),
        ("video", "photo", "Video / photograph"),
        ("photo", "photo", "Photograph"),
        ("screenshot", "screenshot", "Screenshot"),
        ("whatsapp", "message", "WhatsApp / message records"),
        ("email", "email", "Email correspondence"),
        ("agreement", "agreement", "Signed agreement"),
        ("bank statement", "bank_statement", "Bank statement"),
        ("medical report", "medical_report", "Medical report"),
        ("cheque", "cheque", "Cheque / bank memo"),
    ]
    found: list[dict] = []
    seen: set[str] = set()
    for phrase, evidence_type, description in candidates:
        if phrase in lowered and description not in seen:
            seen.add(description)
            found.append(
                {
                    "file_reference": phrase,
                    "evidence_type": evidence_type,
                    "description": f"Citizen says they have: {description}",
                }
            )
    return found


def _infer_incident_type(text: str) -> str:
    lowered = (text or "").lower()
    ordered = [
        "family_domestic_violence",
        "family_maintenance",
        "cyber_crime",
        "rti_refusal",
        "corruption_complaint",
        "debt_recovery",
        "consumer_fraud",
        "tenant_landlord_deposit",
        "wage_theft",
        "wrongful_termination",
        "property_dispute",
        "criminal_assault",
        "criminal_fraud",
    ]
    for incident_type in ordered:
        if any(keyword in lowered for keyword in _INCIDENT_KEYWORDS[incident_type]):
            if incident_type == "consumer_fraud" and "service" in lowered:
                return "consumer_deficiency"
            if incident_type == "wage_theft" and not any(
                word in lowered for word in ("salary", "wage", "deduction", "unpaid")
            ):
                return "wrongful_termination"
            return incident_type
    return "civil"


def _clean_sentences(text: str) -> list[str]:
    raw_parts = re.split(r"(?<=[.!?])\s+|\n+", text or "")
    sentences: list[str] = []
    for part in raw_parts:
        clean = " ".join(part.split()).strip()
        clean = re.sub(r"^User:\s*", "", clean, flags=re.IGNORECASE)
        if len(clean) < 8:
            continue
        if _GREETING_RE.match(clean) and len(clean) < 40:
            continue
        if clean.lower() in {"user", "assistant"}:
            continue
        if clean not in sentences:
            sentences.append(clean)
    return sentences


def _fallback_extract(state: CaseState, ocr_docs: list[dict]) -> dict:
    """Deterministic intake used when Groq is not configured or unavailable."""
    attachment_text = "\n\n".join(
        str(doc.get("extracted_text") or "").strip()
        for doc in ocr_docs
        if str(doc.get("extracted_text") or "").strip()
    )
    full_text = "\n\n".join(
        part for part in (state.raw_narrative or "", attachment_text) if part
    )
    current_text = _last_user_turn(state.raw_narrative or "") or attachment_text
    sentences = _clean_sentences(full_text)
    current_sentences = [
        sentence for sentence in _clean_sentences(current_text)
        if sentence in sentences or not sentences
    ]

    if _GREETING_RE.match(current_text or "") and len(current_text.strip()) < 50:
        return {
            "intent": "GREETING",
            "new_facts": [],
            "follow_up_query": "Please describe the legal incident you need help with.",
            "conversational_response": (
                "Namaste. I am your legal intake assistant. Please describe what "
                "happened, including dates, the other party, and any documents you have."
            ),
            "language_preference": _detect_language(
                current_text, state.language_preference
            ),
            "state_jurisdiction": state.state_jurisdiction,
            "structured_facts": {
                "incident_type": None,
                "case_title": None,
                "incident_summary": None,
                "incident_date": None,
                "parties": [],
                "timeline": [],
                "urgency_level": "standard",
                "monetary_value_inr": None,
                "key_facts": [],
                "missing_information": ["a description of the legal incident"],
                "dispute_context": None,
            },
            "evidence_inventory": [],
        }

    if len(re.findall(r"[a-zA-Z\u0900-\u097F]{3,}", current_text)) == 0:
        return {
            "intent": "ABSURD",
            "new_facts": [],
            "follow_up_query": "Please describe the legal incident in a few sentences.",
            "conversational_response": (
                "I could not understand that input. Please describe the incident, "
                "who was involved, when it happened, and what outcome you want."
            ),
            "language_preference": _detect_language(
                current_text, state.language_preference
            ),
            "state_jurisdiction": state.state_jurisdiction,
            "structured_facts": {
                "incident_type": None,
                "case_title": None,
                "incident_summary": None,
                "incident_date": None,
                "parties": [],
                "timeline": [],
                "urgency_level": "standard",
                "monetary_value_inr": None,
                "key_facts": [],
                "missing_information": ["a description of the legal incident"],
                "dispute_context": None,
            },
            "evidence_inventory": [],
        }

    incident_type = _infer_incident_type(full_text)
    jurisdiction = state.state_jurisdiction or _detect_state(full_text)
    incident_date = _extract_date(full_text)
    amount = _extract_amount(full_text)
    urgency = "immediate" if re.search(
        r"\b(urgent|emergency|immediately|today|right now)\b", full_text, re.I
    ) else "standard"

    summary_source = sentences or ([current_text.strip()] if current_text.strip() else [])
    summary = " ".join(summary_source[:3]).strip()
    if len(summary) > 600:
        summary = summary[:600].rsplit(" ", 1)[0] + "..."

    parties = [
        Party(
            role="complainant",
            description="Person describing the dispute",
        )
    ]
    lowered = full_text.lower()
    for term, label in _RESPONDENT_TERMS:
        if term in lowered:
            parties.append(Party(role="respondent", description=label))
            break

    missing: list[str] = []
    if not incident_date:
        missing.append("the date of the incident")
    if not jurisdiction:
        missing.append("the state where the incident occurred")
    if len(parties) < 2:
        missing.append("the name and contact details of the other party")
    mentioned_evidence = _detect_mentioned_evidence(full_text)
    if not state.uploaded_files and not mentioned_evidence:
        missing.append("documents or other evidence supporting your claim")
    if incident_type in {
        "consumer_deficiency", "consumer_fraud", "debt_recovery",
    } and amount is None:
        missing.append("the amount you are claiming")

    case_title = f"{incident_type.replace('_', ' ').title()} - {jurisdiction or 'India'}"
    evidence_inventory = list(mentioned_evidence)
    if not ocr_docs:
        evidence_inventory.extend(
            {
                "file_reference": filename,
                "evidence_type": "other",
                "description": f"Uploaded file: {os.path.basename(filename)}",
            }
            for filename in state.uploaded_files
        )

    return {
        "intent": "LEGAL",
        "new_facts": current_sentences[:8],
        "follow_up_query": f"Could you provide {missing[0].lower()}?" if missing else "",
        "conversational_response": None,
        "language_preference": _detect_language(
            full_text, state.language_preference
        ),
        "state_jurisdiction": jurisdiction,
        "structured_facts": {
            "incident_type": incident_type,
            "case_title": case_title[:90],
            "incident_summary": summary or "No summary provided.",
            "incident_date": incident_date,
            "parties": [party.model_dump() for party in parties],
            "timeline": [f"Incident date: {incident_date}"] if incident_date else [],
            "urgency_level": urgency,
            "monetary_value_inr": amount,
            "key_facts": (sentences or [current_text.strip()])[:8],
            "missing_information": missing,
            "dispute_context": (
                f"State jurisdiction: {jurisdiction}" if jurisdiction else None
            ),
        },
        "evidence_inventory": evidence_inventory,
    }


# ─────────────────────────────────────────────
# OCR
# ─────────────────────────────────────────────

def _local_text_extract(filename: str) -> Optional[dict]:
    """Extract text from common document types without a paid OCR service.

    Images still need Sarvam OCR or Tesseract. Text files, PDFs and DOCX
    attachments are readable in local fallback mode.
    """
    path = Path(filename)
    suffix = path.suffix.lower()
    text = ""

    try:
        if suffix in {".txt", ".md", ".csv"}:
            text = path.read_text(encoding="utf-8", errors="replace")
        elif suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            text = "\n\n".join(
                page.extract_text() or "" for page in reader.pages[:30]
            )
        elif suffix == ".docx":
            import docx

            document = docx.Document(str(path))
            text = "\n".join(
                paragraph.text for paragraph in document.paragraphs
            )
    except Exception as exc:
        logger.warning("Local text extraction failed for %s: %s", filename, exc)
        return None

    if not text.strip():
        return None

    return {
        "filename": filename,
        "extracted_text": text.strip(),
        "ocr_engine": "local-text",
        "language_code": "auto",
        "job_state": "completed",
        "processed_at": datetime.utcnow().isoformat(),
    }


def _ocr_file(filename: str, sarvam_client: SarvamAI, language_code: str):

    logger.info(f"[Akshar] OCR start: {filename}")

    job = sarvam_client.document_intelligence.create_job(
        language=language_code,
        output_format="md",
    )

    job.upload_file(filename)
    job.start()

    status = job.wait_until_complete()

    with tempfile.NamedTemporaryFile(suffix=".zip", delete=False) as tmp_zip:
        zip_path = tmp_zip.name

    try:

        job.download_output(output_path=zip_path)

        with zipfile.ZipFile(zip_path) as zf:

            md_files = sorted([n for n in zf.namelist() if n.endswith(".md")])

            pages = [
                zf.read(name).decode("utf-8", errors="replace")
                for name in md_files
            ]

            text = "\n\n".join(pages)

            text = re.sub(
                r'!\[.*?\]\(data:image/[^;]+;base64,[^)]+\)',
                "[Embedded Image]",
                text,
            )

    finally:

        os.remove(zip_path)

    return {
        "filename": filename,
        "extracted_text": text,
        "ocr_engine": "sarvam-akshar",
        "language_code": language_code,
        "job_state": status.job_state,
        "processed_at": datetime.utcnow().isoformat(),
    }


# ─────────────────────────────────────────────
# LLM extraction
# ─────────────────────────────────────────────

def _call_groq(narrative: str, ocr_context: str, client: Groq):

    content = narrative

    if ocr_context:
        content += f"\n\n--- Attached document text ---\n{ocr_context}"

    return _call_with_retry(content, INTAKE_SYSTEM_PROMPT, client)


# ─────────────────────────────────────────────
# Builders
# ─────────────────────────────────────────────

def _build_structured_facts(data: dict):

    sf = data.get("structured_facts", {})

    parties = []
    for p in sf.get("parties", []):
        p["role"] = p.get("role") or "unknown"
        parties.append(Party(**p))

    VALID_URGENCY = {"immediate", "standard", "low"}
    raw_urgency = sf.get("urgency_level") or "standard"
    urgency = raw_urgency if raw_urgency in VALID_URGENCY else "standard"

    return StructuredFacts(
        incident_type=sf.get("incident_type"),
        case_title=sf.get("case_title"),
        incident_summary=sf.get("incident_summary"),
        incident_date=sf.get("incident_date"),
        parties=parties,
        timeline=sf.get("timeline", []),
        urgency_level=urgency,
        monetary_value_inr=sf.get("monetary_value_inr"),
        key_facts=sf.get("key_facts", []),
        missing_information=sf.get("missing_information", []),
        dispute_context=sf.get("dispute_context"),
    )


def _build_evidence_inventory(data: dict, ocr_docs):

    items = []

    for ev in data.get("evidence_inventory", []):
        # Type guard: LLM might return a string instead of an object
        if isinstance(ev, str):
            items.append(EvidenceItem(
                file_reference=ev,
                evidence_type="other",
                description=f"User mentioned having: {ev}",
                relevance_score=0.7
            ))
            continue

        items.append(EvidenceItem(
            file_reference=ev.get("file_reference") or "unknown_doc",
            evidence_type=ev.get("evidence_type") or "other",
            date_of_evidence=ev.get("date_of_evidence"),
            relevance_score=float(ev.get("relevance_score") or 0.6),
            description=ev.get("description") or "Document attachment",
            tags=ev.get("tags", []),
        ))

    for doc in ocr_docs:

        items.append(EvidenceItem(
            file_reference=doc["filename"],
            evidence_type="other",
            relevance_score=0.5,
            description=f"Uploaded file: {doc['filename']}",
            ocr_text=doc.get("extracted_text"),
            tags=["uploaded"],
        ))

    return items


# ─────────────────────────────────────────────
# Completeness scoring
# ─────────────────────────────────────────────

def _compute_completeness_score(state: CaseState) -> int:
    """Count how many of the 5 critical intake fields are filled."""
    score = 0
    sf = state.structured_facts
    if sf is None:
        return 0
    if sf.incident_type:
        score += 1
    if sf.incident_summary:
        score += 1
    if sf.parties and len(sf.parties) >= 1:
        score += 1
    if sf.key_facts and len(sf.key_facts) >= 1:
        score += 1
    if state.state_jurisdiction:
        score += 1
    return score


# ─────────────────────────────────────────────
# Follow-up generation
# ─────────────────────────────────────────────

def _generate_followup_questions(narrative, missing_information, groq_client):

    if not missing_information:
        return []

    prompt = f"""
Citizen narrative:
{narrative}

Missing information:
{missing_information}

Convert to natural questions.
Return JSON list only.
"""

    response = groq_client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        max_tokens=MAX_FOLLOWUP_TOKENS,
        messages=[
            {"role": "system", "content": FOLLOWUP_SYSTEM_PROMPT},
            {"role": "user", "content": prompt},
        ],
    )

    raw = response.choices[0].message.content.strip()

    raw = re.sub(r"^```(?:json)?\s*", "", raw)
    raw = re.sub(r"\s*```$", "", raw)

    try:
        return json.loads(raw)[:3]

    except Exception:

        questions = []

        for item in missing_information[:3]:

            questions.append(f"Could you provide {item.lower()}?")

        return questions


def _record_intake_completion(state: CaseState, t0: float, status: str) -> None:
    elapsed = round(time.time() - t0, 2)
    state.agent_trace.append(
        f"[Agent1] {status} {datetime.utcnow().isoformat()} ({elapsed}s)"
    )
    state.case_timeline.append(
        f"{datetime.utcnow().isoformat()} — Intake {status}"
    )


# ─────────────────────────────────────────────
# Main entry
# ─────────────────────────────────────────────

def run(state: CaseState, groq_client: Optional[Groq] = None, sarvam_client: Optional[SarvamAI] = None):

    t0 = time.time()

    groq_key = os.environ.get("GROQ_API_KEY")
    groq = groq_client
    if groq is None and groq_key:
        try:
            groq = Groq(api_key=groq_key)
        except Exception as exc:
            logger.warning("Could not initialise Groq client: %s", exc)

    sarvam_key = os.environ.get("SARVAM_API_KEY")
    sarvam = sarvam_client
    if sarvam is None and sarvam_key:
        try:
            sarvam = SarvamAI(api_subscription_key=sarvam_key)
        except Exception as exc:
            logger.warning("Could not initialise Sarvam client: %s", exc)

    lang_code = _detect_language_code(state.language_preference)
    ocr_docs = []

    if state.uploaded_files and sarvam is None:
        for filename in state.uploaded_files:
            local_doc = _local_text_extract(filename)
            if local_doc:
                ocr_docs.append(local_doc)
            else:
                state.human_review_flags.append(
                    f"{os.path.basename(filename)} attached - image OCR needs SARVAM_API_KEY"
                )
    elif state.uploaded_files and sarvam is not None:
        for filename in state.uploaded_files:
            try:
                ocr_docs.append(_ocr_file(filename, sarvam, lang_code))
            except Exception as exc:
                logger.warning(f"OCR failed for {filename}: {exc}")
                local_doc = _local_text_extract(filename)
                if local_doc:
                    ocr_docs.append(local_doc)
                else:
                    state.human_review_flags.append(f"OCR failed for file: {filename}")

    ocr_context = "\n\n".join(
        f"[{d['filename']}]\n{d.get('extracted_text', '')}" for d in ocr_docs
    )

    fallback_used = False
    if groq is None:
        fallback_used = True
        extracted = _fallback_extract(state, ocr_docs)
        state.human_review_flags.append(
            "Local fallback intake used because GROQ_API_KEY is not configured"
        )
        logger.info("[Agent1] GROQ_API_KEY absent - using deterministic fallback intake")
    else:
        try:
            extracted = _call_groq(state.raw_narrative, ocr_context, groq)
        except Exception as exc:
            fallback_used = True
            logger.warning("[Agent1] Groq intake failed (%s) - using fallback", exc)
            state.human_review_flags.append(
                f"Groq intake unavailable ({exc}); local fallback used"
            )
            extracted = _fallback_extract(state, ocr_docs)

    if not isinstance(extracted, dict):
        extracted = _fallback_extract(state, ocr_docs)
        fallback_used = True

    intent = str(extracted.get("intent") or "LEGAL").upper()
    state.conversational_response = extracted.get("conversational_response")

    # If intent is LEGAL, construct a concise response from the new facts.
    if intent == "LEGAL":
        new_facts = extracted.get("new_facts", []) or []
        query = extracted.get("follow_up_query", "") or ""
        clean_facts = [str(fact) for fact in new_facts if fact and len(str(fact)) > 3]
        if clean_facts:
            fact_list = "\n".join([f"• {fact}" for fact in clean_facts])
            state.conversational_response = f"{fact_list}\n\n{query}".strip()
        else:
            state.conversational_response = query

    if not state.conversational_response:
        if intent == "GREETING":
            state.conversational_response = (
                "Greetings. I am your Legal Intake Assistant. Please provide the "
                "details of the incident you wish to report."
            )
        elif intent == "ABSURD":
            state.conversational_response = (
                "I am unable to process your last input. To assist you, I require "
                "a description of a legal incident, such as a consumer dispute or "
                "property matter. Please provide specific details."
            )
        else:
            missing = (extracted.get("structured_facts") or {}).get(
                "missing_information", []
            ) or []
            if missing:
                state.conversational_response = (
                    "I have recorded the initial facts. To proceed, I specifically "
                    f"require details regarding: {', '.join(missing[:2])}."
                )
            else:
                state.conversational_response = (
                    "The necessary facts have been gathered. I am proceeding with "
                    "the legal analysis."
                )

    language = extracted.get("language_preference")
    if language in {"hindi", "english", "hinglish"}:
        state.language_preference = language

    state.state_jurisdiction = (
        extracted.get("state_jurisdiction") or state.state_jurisdiction
    )

    state.ocr_documents = [
        OcrDocument(
            file_ref=doc.get("filename"),
            type="uploaded",
            extracted_text=doc.get("extracted_text"),
            date_detected=doc.get("processed_at"),
        )
        for doc in ocr_docs
    ]

    state.structured_facts = _build_structured_facts(extracted)

    # Ensure facts are null for non-legal intents to prevent early pipeline trigger
    if intent in ["GREETING", "ABSURD"]:
        state.structured_facts.incident_type = None
        state.structured_facts.incident_summary = None
        state.structured_facts.key_facts = []

    state.evidence_inventory = _build_evidence_inventory(extracted, ocr_docs)

    missing = state.structured_facts.missing_information or []

    # A greeting or unintelligible input must never trigger the full pipeline.
    if intent in ["GREETING", "ABSURD"]:
        state.intake_status = "awaiting_user_response"
        if not state.follow_up_questions:
            state.follow_up_questions = [
                "Could you please describe your legal situation?"
            ]
        _record_intake_completion(state, t0, "awaiting response")
        return state

    completeness = _compute_completeness_score(state)
    logger.info(
        "[Agent1] Completeness score: %s/5 | Missing: %s | fallback=%s",
        completeness,
        missing,
        fallback_used,
    )

    if completeness >= 4:
        state.intake_status = "complete"
        state.follow_up_questions = []
    elif missing:
        top_missing = missing[0]
        state.follow_up_questions = [f"Could you provide {top_missing.lower()}?"]
        state.intake_status = "awaiting_user_response"
        _record_intake_completion(state, t0, "awaiting response")
        return state
    else:
        state.follow_up_questions = [
            "Could you share any more details about what happened?"
        ]
        state.intake_status = "awaiting_user_response"
        _record_intake_completion(state, t0, "awaiting response")
        return state

    _record_intake_completion(state, t0, "completed")

    return state